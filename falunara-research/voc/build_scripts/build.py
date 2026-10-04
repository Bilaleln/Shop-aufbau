import csv, re
from collections import Counter, defaultdict
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from items import *

OUT_X = '/home/user/Shop-aufbau/falunara-research/VOC_Master.xlsx'
OUT_MD = '/home/user/Shop-aufbau/falunara-research/voc/VOC3_Deep_Reddit_VOC_Research.md'
CORPUS = list(csv.DictReader(open('/home/user/Shop-aufbau/falunara-research/voc/VOC2_URL_Corpus.csv', encoding='utf-8')))

def hiratio(rows):
    return f"{round(100 * sum(r['Evidence strength'] == 'high' for r in rows) / max(1, len(rows)))}%"

def conf_txt(rows, override=None):
    n, m, k = nmk(rows)
    c = override or conf(rows)
    return f"{c} ({m} Threads, {k} Communities, high-Evidenz {hiratio(rows)})"

def evid(ids):
    return '\n'.join(q(i) for i in ids)

# ---------- compute all items ----------
def run(items):
    out = []
    for it in items:
        rows = sel(it['f'])
        ex = examples(it, rows)
        out.append(dict(it, rows=rows, exids=ex))
    return out

DS, PS, OS, FS, TS, PAT = map(run, [DESIRES, PROBLEMS, OBJECTIONS, FAILED, TRIGGERS, PATTERNS])
assert not WARN, WARN

# phrase bank term level
PH = []
for pat, lab, cat, note in PHRASES:
    f = RX(pat); rows = sel(f)
    PH.append(dict(name=lab, cat=cat, note=note, rows=rows, exids=pick(rows, 3), pat=pat, f=f))

# isolated tags (<=2 threads)
tagrows = defaultdict(list)
for r in CORE:
    for t in r['tags']: tagrows[t].append(r)
ISO = sorted([(t, rs) for t, rs in tagrows.items() if len({r['Corpus ID'] for r in rs}) <= 2], key=lambda x: x[0])

# ====================== XLSX ======================
HDR = Font(bold=True, color='FFFFFF'); FILL = PatternFill('solid', fgColor='4F6228')
WRAP = Alignment(wrap_text=True, vertical='top')
LINK = Font(color='0563C1', underline='single')
wb = Workbook(); wb.remove(wb.active)

def sheet(name, headers, rows, widths, link_cols=()):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for c in ws[1]:
        c.font = HDR; c.fill = FILL; c.alignment = Alignment(wrap_text=True, vertical='center')
    for row in rows:
        ws.append(row)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
    for col in link_cols:
        for r in range(2, ws.max_row + 1):
            c = ws.cell(r, col); v = c.value
            if isinstance(v, str):
                m = re.search(r'https?://\S+', v)
                if m:
                    c.hyperlink = m.group(0); c.font = LINK
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    return ws

# 00_README
readme = [
 ('Zweck', 'VOC-Master für FALUNARA Botanical Body Oil (U.S.-Frauen 40–65+, Kern 48–58; trockene/crepey wirkende reife Körperhaut). Ergebnis von VOC PROMPT 3 (Synthese). Menopause nur als Kontext. FALUNARA-Inhaltsstoffe sind unbekannt – es werden keine erfunden. Keine Ad-Copy.'),
 ('Datengrundlage', f'{len(D)} verifizierte VOC-Zeilen (Batches A/B/C) aus {len(CORPUS)} Threads (VOC2-Corpus C001–C090; 88 Reddit, 2 Inspire.com). Quelle: voc/VOC3_raw_merged.csv. Jedes Zitat wurde in der Extraktion maschinell gegen den gecachten Thread-Text verifiziert.'),
 ('Tab 01_Source_Corpus', 'Thread-Corpus aus VOC2 (ID, Plattform, Community, Titel, URL, Datum, Thema, Relevanz, Qualität 1–5, Notizen) + Anzahl VOC-Zeilen je Thread, davon high-Evidenz und possible_seeding.'),
 ('Tab 02_Raw_VOC', 'Alle 2.414 Rohzeilen unverändert (Zitat verbatim, Paraphrase deutsch) + Spalte „Normalised pattern tags“ (nach Tag-Mapping unten) + Batch.'),
 ('Tab 03_Desires', 'Wünsche/gewünschte Ergebnisse. Evidence = verbatim Zitate mit VOC-ID; Sources = VOC-IDs + URLs der Beispielzitate; Recurrence = Zeilen/Threads/Communities aller zugeordneten Zeilen (ohne possible_seeding); direkt/abgeleitet; Confidence.'),
 ('Tab 04_Problems_Pains', 'Probleme/Schmerzpunkte mit Originalsprache, typischem Kontext, Quellen, Rekurrenz.'),
 ('Tab 05_Objections', 'Einwände mit exakter Sprache, Quelle, Rekurrenz und Forschungs-Implikation (keine Copy).'),
 ('Tab 06_Failed_Solutions', 'Erfolglose Lösungsversuche, Beschwerde/Ergebnis, Quelle, Rekurrenz.'),
 ('Tab 07_Triggers', 'Situationen/Auslöser mit Zitat-Evidenz, Quelle, Rekurrenz.'),
 ('Tab 08_Phrase_Bank', 'Teil A: Begriffs-/Metaphernfamilien mit Regex-Zählung im Zitattext (Zeilen/Threads/Communities). Teil B: alle Zeilen der Kategorie „Exact phrases / metaphors / repeated wording“ (ohne Seeding) als Einzelphrasen.'),
 ('Tab 09_Patterns', 'Muster-Synthese: wiederkehrende Muster (P01–P25) mit allen stützenden VOC-IDs, Gegenevidenz, Konfidenz; danach isolierte Aussagen (Tags in ≤2 Threads).'),
 ('Spalte Recurrence', 'Format „N rows / M threads / K communities“, berechnet mit Python über die zugeordneten Zeilen, OHNE Zeilen mit Flag possible_seeding. Eine Zeile kann mehreren Items zugeordnet sein (Mehrfachzählung über Items hinweg).'),
 ('Spalte Recurrence label', 'stark wiederkehrend = ≥20 Threads; wiederkehrend = 8–19 Threads; begrenzt wiederkehrend = 3–7 Threads; isoliert = ≤2 Threads. Basis: Anzahl distinkter Threads (Corpus IDs).'),
 ('Spalte Confidence', 'hoch = ≥15 Threads UND ≥4 Communities; mittel = ≥6 Threads UND ≥2 Communities; sonst niedrig. In Klammern: Threads, Communities, Anteil high-Evidenz. Konfidenz bezieht sich auf die Wiederkehr im Corpus, NICHT auf Repräsentativität für die U.S.-Gesamtpopulation.'),
 ('Spalte Sensitivity', 'Anzahl ausgeschlossener Seeding-Zeilen; Zählung ohne Zeilen mit non_us_signal/under_40_signal; Anzahl Zeilen mit non_US, U40 und hrt_medical_context; Anteil high-Evidenz.'),
 ('Spalte All supporting VOC IDs', 'Vollständige Liste der zugeordneten VOC-IDs (Audit-Pfad zu 02_Raw_VOC). Zuordnung erfolgt über normalisierte Tags, ggf. kombiniert mit Kategorie- und Regex-Filtern auf den Zitattext.'),
 ('Hyperlinks', 'In Zellen mit mehreren URLs ist die erste URL verlinkt; alle URLs stehen im Zelltext.'),
 ('Evidence strength', 'high = explizit, Ich-Perspektive, spezifisch, Körperhaut, Zielalter/-community; medium = Ich-Perspektive aber vage, oder Alter/Körperzone unklar; low = Second-hand, Off-target-Körperzone oder wahrscheinlich non-US/unter 40.'),
 ('Flag possible_seeding', f'{sum("possible_seeding" in r["flags"] for r in D)} Zeilen; Werbe-/Brand-Ton, Links, Influencer-Accounts, C031-OP. In allen Zählungen der Tabs 03–09 AUSGESCHLOSSEN (bleiben in 02_Raw_VOC sichtbar).'),
 ('Flag non_us_signal', f'{sum("non_us_signal" in r["flags"] for r in D)} Zeilen mit Hinweis auf Nicht-U.S.-Herkunft (Kanada, UK, AU, NZ, EU). Mitgezählt; Sensitivität separat ausgewiesen.'),
 ('Flag under_40_signal', f'{sum("under_40_signal" in r["flags"] for r in D)} Zeilen von Posterinnen unter 40. Mitgezählt; Sensitivität separat ausgewiesen.'),
 ('Flag hrt_medical_context', f'{sum("hrt_medical_context" in r["flags"] for r in D)} Zeilen mit HRT-/medizinischem Kontext. Nur Kontext – keine Produktclaims ableiten.'),
 ('Flag face_only_transferable', f'{sum("face_only_transferable" in r["flags"] for r in D)} Zeilen zu Gesicht, nur übertragbar.'),
 ('Flag possible_male_poster', f'{sum("possible_male_poster" in r["flags"] for r in D)} Zeilen (C047-OP erwähnt „my wife“).'),
 ('Bekannter Corpus-Fehler', 'C034: Laut VOC2 ist die OP 55 – im Text stammt „55“ von einer anderen Posterin (Eintrag 26). OP-Zeilen haben daher kein Alter.'),
 ('Limitationen', 'Reddit-dominiert (88/90 Threads), U.S.-Wohnsitz nicht verifizierbar, Alter nur in 438 Zeilen angegeben, max. ~100 Kommentare je Thread im Cache, Tagging durch drei Extraktions-Batches (normalisiert, aber Restunschärfe), Zählungen = Erwähnungen im Corpus, nicht Marktanteile. Siehe MD-Report §11.'),
 ('Tag-Normalisierung', 'Mapping der Batch-spezifischen new_-Tags auf normalisierte Tags (unten). Nicht gemappte new_-Tags: Präfix „new_“ entfernt. question_oil_or_lotion → oil_vs_lotion_confusion; new_seeding_suspected → marketing_distrust.'),
]
for k, v in sorted(MAP.items(), key=lambda x: (x[1], x[0])):
    readme.append((f'Mapping: {k}', v))
sheet('00_README', ['Thema', 'Erläuterung'], readme, [34, 140])

# 01_Source_Corpus
cnt = Counter(r['Corpus ID'] for r in D)
hic = Counter(r['Corpus ID'] for r in D if r['Evidence strength'] == 'high')
sdc = Counter(r['Corpus ID'] for r in D if 'possible_seeding' in r['flags'])
rows = []
for c in CORPUS:
    note = c['Notes']
    if c['ID'] == 'C034':
        note += ' | VOC3-Korrektur: OP-Alter 55 nicht belegt (55 stammt von anderer Posterin, Eintrag 26).'
    rows.append([c['ID'], c['Platform'], c['Community'], c['Thread title'], c['URL'], c['Date if available'], c['Topic'],
                 c['Why relevant'], c['Expected VOC categories'], int(c['Quality score']) if c['Quality score'].isdigit() else c['Quality score'],
                 note, cnt[c['ID']], hic[c['ID']], sdc[c['ID']]])
sheet('01_Source_Corpus', ['Corpus ID', 'Platform', 'Community', 'Title', 'URL', 'Date', 'Topic', 'Relevance', 'Expected VOC categories',
                           'Quality (1-5)', 'Notes', 'VOC rows', 'VOC rows high', 'VOC rows possible_seeding'],
      rows, [10, 12, 22, 45, 50, 12, 30, 45, 35, 10, 50, 10, 10, 12], link_cols=(5,))

# 02_Raw_VOC
rows = [[r['VOC ID'], r['Category'], r['Exact short quote'], r['Paraphrased meaning'], r['Source URL'], r['Community'], r['Context'],
         r['Recurrence/pattern tag'], ';'.join(r['tags']), r['Evidence strength'], r['Thread URL'], r['Corpus ID'],
         r['Potential research implication'], int(r['Stated age']) if r['Stated age'].isdigit() else r['Stated age'], r['Body area'], r['Flags'],
         r['VOC ID'][0]] for r in D]
sheet('02_Raw_VOC', ['VOC ID', 'Category', 'Exact short quote', 'Paraphrased meaning (DE)', 'Source URL', 'Community', 'Context',
                     'Pattern tag (original)', 'Normalised pattern tags', 'Evidence strength', 'Thread URL', 'Corpus ID',
                     'Potential research implication', 'Stated age', 'Body area', 'Flags', 'Batch'],
      rows, [9, 22, 60, 45, 40, 18, 35, 28, 28, 10, 40, 9, 45, 8, 12, 22, 6], link_cols=(5, 11))

def common(it):
    return [rec(it['rows']), label(it['rows'])]

# 03
rows = [[n + 1, it['name'], evid(it['exids']), srcs(it['exids']), *common(it), it['kind'], conf_txt(it['rows']),
         sens(it['f']), it['note'], ids_str(it['rows'])] for n, it in enumerate(DS)]
sheet('03_Desires', ['#', 'Desire', 'Evidence (verbatim)', 'Sources (VOC ID: URL)', 'Recurrence', 'Recurrence label', 'Direct/inferred',
                     'Confidence', 'Sensitivity', 'Notes', 'All supporting VOC IDs'], rows,
      [5, 40, 70, 55, 30, 18, 14, 30, 45, 40, 60], link_cols=(4,))
# 04
rows = [[n + 1, it['name'], evid(it['exids']), it['ctx'], srcs(it['exids']), *common(it), conf_txt(it['rows']), sens(it['f']),
         ids_str(it['rows'])] for n, it in enumerate(PS)]
sheet('04_Problems_Pains', ['#', 'Problem/pain', 'Language (verbatim)', 'Context', 'Sources (VOC ID: URL)', 'Recurrence', 'Recurrence label',
                            'Confidence', 'Sensitivity', 'All supporting VOC IDs'], rows, [5, 40, 70, 40, 55, 30, 18, 30, 45, 60], link_cols=(5,))
# 05
rows = [[n + 1, it['name'], evid(it['exids']), srcs(it['exids']), *common(it), it['impl'], conf_txt(it['rows']), sens(it['f']),
         ids_str(it['rows'])] for n, it in enumerate(OS)]
sheet('05_Objections', ['#', 'Objection', 'Exact language (verbatim)', 'Source (VOC ID: URL)', 'Recurrence', 'Recurrence label', 'Implication (research, no copy)',
                        'Confidence', 'Sensitivity', 'All supporting VOC IDs'], rows, [5, 40, 70, 55, 30, 18, 50, 30, 45, 60], link_cols=(4,))
# 06
rows = [[n + 1, it['name'], it['out'], evid(it['exids']), srcs(it['exids']), *common(it), conf_txt(it['rows']), sens(it['f']),
         ids_str(it['rows'])] for n, it in enumerate(FS)]
sheet('06_Failed_Solutions', ['#', 'Attempted solution', 'Complaint/outcome (Zusammenfassung)', 'Complaint/outcome (verbatim)', 'Source (VOC ID: URL)',
                              'Recurrence', 'Recurrence label', 'Confidence', 'Sensitivity', 'All supporting VOC IDs'], rows,
      [5, 38, 45, 70, 55, 30, 18, 30, 45, 60], link_cols=(5,))
# 07
rows = [[n + 1, it['name'], evid(it['exids']), srcs(it['exids']), *common(it), conf_txt(it['rows']), sens(it['f']),
         ids_str(it['rows'])] for n, it in enumerate(TS)]
sheet('07_Triggers', ['#', 'Situation/trigger', 'Quote/evidence (verbatim)', 'Source (VOC ID: URL)', 'Recurrence', 'Recurrence label',
                      'Confidence', 'Sensitivity', 'All supporting VOC IDs'], rows, [5, 40, 70, 55, 30, 18, 30, 45, 60], link_cols=(4,))
# 08
rows = []
for it in PH:
    rows.append(['A: Begriffsfamilie', it['name'], it['cat'], srcs(it['exids']), evid(it['exids']),
                 (it['note'] + ' ' if it['note'] else '') + f"Regex auf Zitattext: {it['pat']}", rec(it['rows']), label(it['rows']), ids_str(it['rows'])])
tr = {t: rs for t, rs in tagrows.items()}
for r in CORE:
    if r['Category'] != PHR: continue
    tagrec = '; '.join(f"{t}: {rec(tr[t])}" for t in r['tags'])
    rows.append(['B: Einzelphrase', r['Exact short quote'], 'Exact phrases / metaphors / repeated wording (' + ';'.join(r['tags']) + ')',
                 f"{r['VOC ID']}: {r['Source URL']}", q(r['VOC ID']), f"DE: {r['Paraphrased meaning']} | Flags: {r['Flags'] or '–'} | Alter: {r['Stated age'] or '–'}",
                 '1 row (Tag-Rekurrenz: ' + tagrec + ')', '', r['VOC ID']])
sheet('08_Phrase_Bank', ['Teil', 'Phrase/metaphor', 'Category', 'Source (VOC ID: URL)', 'Example quotes (verbatim)', 'Notes', 'Recurrence',
                         'Recurrence label', 'All supporting VOC IDs'], rows, [16, 45, 30, 55, 70, 50, 40, 18, 50], link_cols=(4,))
# 09
rows = []
for it in PAT:
    cnt_txt = '\n'.join(q(i) for i in it['counter']) + ('\n' if it['counter'] else '') + it['cnote']
    rows.append([it['key'], it['name'], ids_str(it['rows']), evid(it['exids']), srcs(it['exids']), cnt_txt, conf_txt(it['rows']),
                 rec(it['rows']), label(it['rows']), sens(it['f'])])
for t, rs in ISO:
    rows.append(['ISO', f'Isolierte Aussage – Tag „{t}“', ids_str(rs), evid([r['VOC ID'] for r in rs[:2]]), srcs([r['VOC ID'] for r in rs[:2]]),
                 '– (zu geringe Basis)', 'niedrig (isoliert)', rec(rs), label(rs), ''])
sheet('09_Patterns', ['Pattern ID', 'Pattern', 'Supporting VOC IDs (all)', 'Key quotes (verbatim)', 'Sources (VOC ID: URL)',
                      'Counter-evidence', 'Confidence', 'Recurrence', 'Recurrence label', 'Sensitivity'], rows,
      [9, 50, 60, 70, 55, 60, 30, 30, 18, 45], link_cols=(5,))

wb.save(OUT_X)

# verify
wb2 = load_workbook(OUT_X)
print('Sheets:', wb2.sheetnames)
for ws in wb2.worksheets:
    print(f'  {ws.title}: {ws.max_row - 1} data rows, {ws.max_column} cols')
# verbatim check: every quote string in 03-09 must be from CSV
allq = {r['VOC ID']: r['Exact short quote'] for r in D}
bad = 0
for ws in wb2.worksheets[3:]:
    for row in ws.iter_rows(min_row=2, values_only=True):
        for v in row:
            if isinstance(v, str):
                for m in re.finditer(r'„(.*?)“ \(([ABC]\d{4})\)', v, re.S):
                    if allq.get(m.group(2)) != m.group(1): bad += 1
print('quote mismatches:', bad)

